"""Safe text extraction from uploaded PDF / DOCX / TXT files.

Validation is by content (magic bytes), not by the client-supplied filename.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass


class UnsupportedFile(ValueError):
    pass


@dataclass
class ExtractedText:
    text: str
    file_type: str
    pages: int | None = None
    warnings: list[str] | None = None


def detect_type(data: bytes, filename: str = "") -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        raise UnsupportedFile("ZIP container is not a DOCX document")
    if b"\x00" in data[:4096]:
        raise UnsupportedFile("Binary file is not a supported format (PDF, DOCX, TXT)")
    for enc in ("utf-8", "utf-16"):
        try:
            data[:65536].decode(enc)
            return "txt"
        except UnicodeDecodeError:
            continue
    try:
        data[:65536].decode("latin-1")
        return "txt"
    except UnicodeDecodeError:
        raise UnsupportedFile("Unrecognised file encoding")


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _fragmented(text: str) -> bool:
    """True when nearly every non-empty line holds a single word (broken text-object ordering)."""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    words = len(text.split())
    return words >= 40 and words / max(len(lines), 1) < 1.5


def extract_pdf(data: bytes) -> ExtractedText:
    from pypdf import PdfReader

    warnings = []
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise UnsupportedFile("Encrypted PDF")
    if len(reader.pages) > 50:
        raise UnsupportedFile("PDF has more than 50 pages")
    pages = [p.extract_text() or "" for p in reader.pages]
    text = "\n".join(pages)
    # pypdf emits one word per line for some producers (e.g. Google Docs' Skia renderer), which destroys the
    # line/section structure the parser relies on. pdfplumber groups words into visual lines.
    if len(text.strip()) < 50 or _fragmented(text):
        try:  # fall back to pdfplumber for unusual layouts
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                alt = "\n".join((p.extract_text() or "") for p in pdf.pages)
            if len("".join(alt.split())) >= len("".join(text.split())) * 0.8:
                text = alt
        except Exception:
            warnings.append("pdfplumber fallback failed")
    if len(text.strip()) < 50:
        warnings.append("Very little text extracted — the PDF may be scanned images (OCR is not supported).")
    return ExtractedText(_clean(text), "pdf", len(reader.pages), warnings)


def extract_docx(data: bytes) -> ExtractedText:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return ExtractedText(_clean("\n".join(parts)), "docx", None, [])


def extract_txt(data: bytes) -> ExtractedText:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return ExtractedText(_clean(data.decode(enc)), "txt", None, [])
        except UnicodeDecodeError:
            continue
    raise UnsupportedFile("Unrecognised text encoding")


def extract_text(data: bytes, filename: str = "") -> ExtractedText:
    kind = detect_type(data, filename)
    return {"pdf": extract_pdf, "docx": extract_docx, "txt": extract_txt}[kind](data)
